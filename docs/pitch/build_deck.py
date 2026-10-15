"""Builds Temple_Platform_Demo.pptx from screenshots in shots/. Run: python build_deck.py"""
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

MAROON, SAFFRON, CREAM, INK, GOLD = (RGBColor(*c) for c in
    [(0x7B, 0x1A, 0x1A), (0xC8, 0x7A, 0x1E), (0xFA, 0xF5, 0xEA), (0x2B, 0x20, 0x20), (0xE8, 0xB2, 0x3A)])
SH = "shots/"
LOGO = "../../frontend/public/logo.png"


def clean(name, box, sample):
    """Paint over the Next.js dev badge so the demo looks like production."""
    im = Image.open(SH + name + ".png").convert("RGB")
    ImageDraw.Draw(im).rectangle(box, fill=im.getpixel(sample))
    im.save(SH + name + "_c.png")
    return SH + name + "_c.png"


imgs = {
    "home": clean("home", (10, 835, 150, 890), (160, 862)),
    "telugu": clean("telugu", (10, 835, 150, 890), (160, 862)),
    "donations": clean("donations", (10, 835, 150, 890), (160, 862)),
    "mobile": clean("mobile", (10, 780, 150, 830), (5, 805)),
    "login": clean("login", (10, 835, 80, 890), (200, 862)),
}

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill
    if line: r.line.color.rgb = line; r.line.width = Pt(1)
    else: r.line.fill.background()
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, t, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
        p.space_after = Pt(6)
    return tb


def slide(title, kicker=None):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = CREAM
    rect(s, 0, 0, 13.333, 0.12, SAFFRON)
    text(s, 0.6, 0.35, 12, 0.9, title, 32, MAROON, True, font="Georgia")
    if kicker: text(s, 0.6, 1.05, 12, 0.5, kicker, 16, SAFFRON)
    text(s, 0.6, 7.05, 8, 0.3, "Temple Platform  |  Vriksha Studio", 10, RGBColor(0x88, 0x77, 0x77))
    return s


def shot(s, key, x, y, w):
    """Image with a thin browser-style frame."""
    im = Image.open(imgs[key]); h = w * im.height / im.width
    rect(s, x - 0.05, y - 0.05, w + 0.1, h + 0.1, RGBColor(255, 255, 255), RGBColor(0xD8, 0xC8, 0xB0))
    s.shapes.add_picture(imgs[key], Inches(x), Inches(y), Inches(w))
    return h


def bullets(s, x, y, w, items, size=18):
    for i, it in enumerate(items):
        rect(s, x, y + i * 0.62 + 0.12, 0.12, 0.12, SAFFRON, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.3, y + i * 0.62, w - 0.3, 0.6, it, size)


def card(s, x, y, w, h, head, body, accent=MAROON):
    rect(s, x, y, w, h, RGBColor(255, 255, 255), RGBColor(0xE6, 0xD9, 0xC3))
    rect(s, x, y, w, 0.1, accent)
    text(s, x + 0.2, y + 0.2, w - 0.4, 0.5, head, 18, accent, True, font="Georgia")
    text(s, x + 0.2, y + 0.75, w - 0.4, h - 0.9, body, 14)


# 1 Title
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = MAROON
rect(s, 0, 6.9, 13.333, 0.6, SAFFRON)
s.shapes.add_picture(LOGO, Inches(0.9), Inches(1.6), Inches(2.4))
text(s, 3.8, 1.5, 9, 0.5, "ఓం గం గణపతయే నమః", 22, GOLD)
text(s, 3.8, 2.1, 9, 1.8, "A Digital Home for Every Temple", 46, RGBColor(255, 255, 255), True, font="Georgia")
text(s, 3.8, 3.9, 9, 1, "Multilingual website, seva booking and temple management, proven at Sri Varasidhi Vinayaka Swamy Devasthanam, Thorur",
     20, RGBColor(0xF3, 0xE3, 0xC8))
text(s, 3.8, 5.4, 9, 0.5, "Live demo  |  Packages  |  Budget", 18, GOLD, True)
text(s, 0.9, 6.95, 11, 0.4, "Vriksha Studio  |  October 2026", 14, MAROON, True)

# 2 Why
s = slide("Why temples need more than a website", "What devotees and temple offices struggle with today")
card(s, 0.6, 1.8, 3.9, 3.6, "For devotees", ["Timings and festival dates hard to find", "Seva booking by phone or in person", "Information only in one language", "Poor experience on mobile"])
card(s, 4.7, 1.8, 3.9, 3.6, "For the temple office", ["Paper receipts and registers", "No single record of donors", "Accounts scattered in notebooks", "No trail of who changed what"], SAFFRON)
card(s, 8.8, 1.8, 3.9, 3.6, "For trustees", ["Little visibility into income", "Trust depends on one person", "Hard to show transparency", "Data privacy obligations (DPDPA)"])
text(s, 0.6, 5.8, 12, 1, "One platform answers all three: a public site, an admin office and an audit trail.", 22, MAROON, True, font="Georgia")

# 3 Home
s = slide("Live demo: the public website", "Running on localhost, the same codebase that serves svvdthorur.org")
shot(s, "home", 0.6, 1.7, 7.9)
bullets(s, 8.9, 1.9, 4.2, ["Deity carousel with Sanskrit/Telugu greeting", "Darshan timings always on top", "Announcements feed", "Clear navigation: Timings, Poojas, Festivals, Gallery, About, Contact", "Temple branding in one theme"], 16)

# 4 Multilingual + mobile
s = slide("Five languages, built for the phone", "English, Telugu, Tamil, Kannada, Hindi: each on its own URL")
shot(s, "telugu", 0.6, 1.7, 7.0)
im = Image.open(imgs["mobile"]); mh = 4.9; mw = mh * im.width / im.height
rect(s, 8.25, 1.65, mw + 0.1, mh + 0.1, RGBColor(255, 255, 255), RGBColor(0xD8, 0xC8, 0xB0))
s.shapes.add_picture(imgs["mobile"], Inches(8.3), Inches(1.7), Inches(mw), Inches(mh))
bullets(s, 10.9, 1.9, 2.3, ["Per-language SEO", "Mobile-first", "Add a language with one file"], 14)
text(s, 0.6, 5.3, 7, 1.2, "Search engines see every language: canonical and hreflang tags, sitemap, structured data.", 16, INK)

# 5 Devotee journey
s = slide("Devotee journey: book a seva in minutes", "No account needed; account optional for history")
shot(s, "login", 0.6, 1.7, 5.4)
steps = [("1", "Choose a seva"), ("2", "Verify email with OTP"), ("3", "Booking confirmed + QR ticket by email"), ("4", "Pay at the counter (or online with gateway add-on)"), ("5", "Staff scans QR at the temple")]
for i, (n, t) in enumerate(steps):
    y = 1.8 + i * 0.95
    rect(s, 6.6, y, 0.6, 0.6, MAROON, shape=MSO_SHAPE.OVAL); text(s, 6.6, y + 0.05, 0.6, 0.5, n, 20, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
    text(s, 7.4, y + 0.05, 5.5, 0.6, t, 18)
text(s, 6.6, 6.5, 6.5, 0.5, "Optional sign-in shows 'My Bookings'; personal blessing email on birthdays and anniversaries.", 13, SAFFRON)

# 6 Donations transparency
s = slide("Donations with transparency", "Archive of historical donors today, online giving as an add-on")
shot(s, "donations", 0.6, 1.7, 7.4)
bullets(s, 8.4, 1.9, 4.7, ["Historic donors, land and item donors, committee archive", "Admin records donations and issues PDF receipts", "A receipt, once issued, cannot be edited or deleted", "Every rupee posts to the finance ledger", "Gateway (UPI/cards) available as add-on"], 16)

# 7 Admin
s = slide("Behind the scenes: the admin panel", "Staff manage everything without a developer: 14 modules, 5 roles")
mods = ["Dashboard", "Announcements", "Festivals", "Poojas & Sevas", "Gallery", "Timings", "Seva Tickets + QR scan", "Donors & Donations", "Finance ledger", "Members", "Messages", "Users", "Audit log", "Temple profile"]
for i, m in enumerate(mods):
    c, r = i % 4, i // 4
    rect(s, 0.6 + c * 3.1, 1.9 + r * 0.85, 2.9, 0.65, RGBColor(255, 255, 255), RGBColor(0xE6, 0xD9, 0xC3))
    text(s, 0.7 + c * 3.1, 1.95 + r * 0.85, 2.7, 0.55, m, 15, MAROON, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 5.25, 12, 0.5, "Roles", 20, SAFFRON, True, font="Georgia")
roles = [("Super Admin", "Everything + users + audit"), ("Admin", "All content, donors, finance"), ("Trustee", "Read-only oversight"), ("Staff", "Content, tickets, messages"), ("Devotee", "Own bookings only")]
for i, (r_, d) in enumerate(roles):
    x = 0.6 + i * 2.5
    rect(s, x, 5.8, 2.35, 1.1, MAROON if i == 0 else RGBColor(255, 255, 255), RGBColor(0xE6, 0xD9, 0xC3))
    col = RGBColor(255, 255, 255) if i == 0 else MAROON
    text(s, x + 0.1, 5.85, 2.15, 0.4, r_, 15, col, True)
    text(s, x + 0.1, 6.25, 2.15, 0.6, d, 12, RGBColor(0xF3, 0xE3, 0xC8) if i == 0 else INK)

# 8 Architecture
s = slide("How it is built", "Modern, tested, containerised: low cost to run, easy to hand over")
boxes = [("Devotee / Staff\n(browser)", 0.6, RGBColor(255, 255, 255), MAROON), ("Nginx + TLS\n(Let's Encrypt)", 3.0, RGBColor(255, 255, 255), MAROON),
         ("Next.js 15 + React 19\npublic site, admin UI", 5.4, MAROON, RGBColor(255, 255, 255)), ("FastAPI (Python)\n52 REST endpoints, RBAC", 8.0, MAROON, RGBColor(255, 255, 255)),
         ("PostgreSQL", 10.7, SAFFRON, RGBColor(255, 255, 255))]
for t, x, f, c in boxes:
    rect(s, x, 2.2, 2.2, 1.3, f, MAROON, MSO_SHAPE.ROUNDED_RECTANGLE); text(s, x, 2.2, 2.2, 1.3, t.split("\n"), 13, c, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if x < 10: text(s, x + 2.15, 2.5, 0.4, 0.5, ">", 22, SAFFRON, True)
for t, x in [("AWS S3\ngallery, tickets, receipts", 6.3), ("AWS SES\nOTP, confirmations, alerts", 9.0)]:
    rect(s, x, 4.3, 2.5, 1.0, RGBColor(255, 255, 255), SAFFRON, MSO_SHAPE.ROUNDED_RECTANGLE); text(s, x, 4.3, 2.5, 1.0, t.split("\n"), 12, INK, True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.1, 5.4, 2.8, ["One server per temple: ~₹1,000 a month", "Docker Compose; GitHub Actions deploys after tests pass", "21 versioned DB migrations, ~40 test modules", "Each temple isolated: own domain, data and branding"], 15)

# 9 Security
s = slide("Trust: security and privacy built in")
items = [("Access control", "Five roles; permissions checked on the server for every action"), ("Two-factor login", "Email OTP; 30-minute idle timeout"), ("Audit trail", "Every change recorded; alerts on bulk deletions"),
         ("Abuse protection", "Rate limiting and bot challenge on public forms"), ("Privacy (DPDPA)", "Privacy policy on site, emails and tickets"), ("Safe by default", "HTTPS, hashed passwords, no AWS keys in app, dependency audits")]
for i, (h, b) in enumerate(items):
    card(s, 0.6 + (i % 3) * 4.15, 1.6 + (i // 3) * 2.7, 3.95, 2.4, h, b, MAROON if i % 2 == 0 else SAFFRON)

# 10 Packages
s = slide("Three packages", "Same platform; choose what the temple needs today and grow later")
pk = [("Essential", "₹55,000", "3-4 weeks", ["Public website, 2 languages", "Admin CMS and roles", "Contact form, SEO", "Hosting setup, training"], SAFFRON),
      ("Standard", "₹1,40,000", "6-8 weeks", ["Everything in Essential", "5 languages", "Seva booking, OTP, QR ticket", "Devotee accounts", "Donors and PDF receipts"], MAROON),
      ("Complete", "₹2,60,000", "10-12 weeks", ["Everything in Standard", "Finance ledger", "Member directory", "Audit log, 2FA, QR scan", "Occasion blessings"], SAFFRON)]
for i, (n, p, t, its, col) in enumerate(pk):
    x = 0.6 + i * 4.15
    rect(s, x, 1.7, 3.95, 5.0, RGBColor(255, 255, 255), col if i != 1 else MAROON)
    rect(s, x, 1.7, 3.95, 1.5, col)
    text(s, x, 1.75, 3.95, 0.5, n, 22, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER, "Georgia")
    text(s, x, 2.2, 3.95, 0.6, p, 30, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
    text(s, x, 2.8, 3.95, 0.4, t + "  |  one-time, excl. GST", 12, RGBColor(255, 255, 255), False, PP_ALIGN.CENTER)
    text(s, x + 0.3, 3.4, 3.4, 3.2, ["- " + it for it in its], 15)

# 11 Running cost
s = slide("What it costs to run, and the multi-temple saving")
rows = [["", "Essential", "Standard", "Complete"], ["Hosting, email, domain, backups / yr", "₹13,600", "₹15,900", "₹23,700"], ["Maintenance & support / yr", "₹24,000", "₹48,000", "₹84,000"],
        ["Recurring / yr", "₹37,600", "₹63,900", "₹1,07,700"], ["Year 1 total", "₹92,600", "₹2,03,900", "₹3,67,700"], ["3-year total", "₹1,67,800", "₹3,31,700", "₹5,83,100"]]
t = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.6), Inches(7.6), Inches(3.4)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = t.cell(r, c); cell.text = v or " "; p = cell.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(14); p.runs[0].font.name = "Calibri"
        if c: p.alignment = PP_ALIGN.RIGHT
        cell.fill.solid(); cell.fill.fore_color.rgb = MAROON if r == 0 else (RGBColor(255, 255, 255) if r % 2 else RGBColor(0xF5, 0xEF, 0xE6))
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255) if r == 0 else INK; p.runs[0].font.bold = (r == 0 or r >= 4)
card(s, 8.6, 1.6, 4.1, 3.4, "5 temples, Standard", ["One-time platform work: ₹90,000", "Temples 2-5: 25% off build", "Year 1 total: ₹9,69,500", "Without programme: ₹10,19,500"], SAFFRON)
text(s, 0.6, 5.4, 12, 1.4, ["Estimates in INR excl. 18% GST; online payment gateway is an optional ₹35,000 add-on.", "Hosting alone is about ₹800-1,200 a month per temple. Full detail in the Budget document."], 14, RGBColor(0x55, 0x44, 0x44))

# 12 Next steps
s = slide("Next steps", "From this demo to a live temple site in weeks")
for i, (n, t) in enumerate([("Week 0", "Choose package, sign quotation"), ("Week 1", "Collect content: timings, sevas, photos, committee"), ("Weeks 2-N", "Build on staging; review with temple"),
                            ("Go-live", "Domain, HTTPS, training, handover"), ("After", "30-day hypercare, then maintenance")]):
    x = 0.6 + i * 2.5
    rect(s, x, 2.3, 2.3, 0.7, MAROON if i % 2 == 0 else SAFFRON, shape=MSO_SHAPE.PENTAGON)
    text(s, x, 2.3, 2.1, 0.7, n, 16, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, 3.2, 2.3, 1.5, t, 14)
text(s, 0.6, 5.2, 12, 1, "Pilot offer: the next temple's staging site can be set up for review before any commitment.", 20, MAROON, True, font="Georgia")
text(s, 0.6, 6.2, 12, 0.5, "Chaithanya M  |  vrikshastudio9@gmail.com", 16, SAFFRON)

prs.save("Temple_Platform_Demo.pptx"); print("wrote Temple_Platform_Demo.pptx")
